#!/usr/bin/env python3
import argparse, csv, json, os, re, sys


def norm(v):
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v)).strip()


def rows_from_csv(path):
    with open(path, newline='', encoding='utf-8-sig') as f:
        for row in csv.reader(f):
            yield row


def rows_from_xlsx(path):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise SystemExit("openpyxl is required for XLSX input")
    wb = load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            yield list(row)


def detect_pairs(rows):
    pairs = []
    header_seen = False
    for row in rows:
        vals = [norm(x) for x in row]
        if not any(vals):
            continue
        low = [x.lower() for x in vals]
        if not header_seen and any(x in {"vendor", "вендор", "производитель"} for x in low) and any(x in {"product", "продукт"} for x in low):
            header_seen = True
            continue
        if len(vals) >= 2 and vals[0] and vals[1]:
            pairs.append((vals[0], vals[1]))
        elif len(vals) == 1 and vals[0]:
            pairs.append((vals[0], ""))
    return pairs


def main():
    ap = argparse.ArgumentParser(description="Normalize a vendor/product catalog from CSV or XLSX")
    ap.add_argument("input")
    ap.add_argument("--exclude", action="append", default=[])
    ap.add_argument("--json", dest="json_out")
    args = ap.parse_args()
    ext = os.path.splitext(args.input)[1].lower()
    rows = rows_from_xlsx(args.input) if ext == '.xlsx' else rows_from_csv(args.input)
    pairs = detect_pairs(rows)
    excluded = {norm(x).casefold() for x in args.exclude}
    out = []
    seen = set()
    for vendor, product in pairs:
        if vendor.casefold() in excluded or product.casefold() in excluded:
            continue
        key = (vendor.casefold(), product.casefold())
        if key in seen:
            continue
        seen.add(key)
        out.append({"vendor": vendor, "product": product})
    result = {
        "vendors": sorted({x["vendor"] for x in out}),
        "products": [x for x in out if x["product"]],
        "counts": {
            "vendors": len({x["vendor"] for x in out}),
            "products": sum(1 for x in out if x["product"])
        }
    }
    text = json.dumps(result, ensure_ascii=False, indent=2)
    if args.json_out:
        with open(args.json_out, 'w', encoding='utf-8') as f:
            f.write(text + "\n")
    else:
        print(text)

if __name__ == '__main__':
    main()
